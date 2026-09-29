"""inspect_v82_sheets2.py - layout + merged-cell evidence for the v8.2 workbook (read-only)."""
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(BASE, "Global-Smile_2026-v8.2-final-review.xlsm")
wb = openpyxl.load_workbook(WB, keep_vba=True, data_only=True)


def cl(i):
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def dump(name, rows, cols, label=None):
    ws = wb[name]
    print("=" * 110)
    print("SHEET %s  dims=%s  max_row=%s max_col=%s" % (name, ws.dimensions, ws.max_row, ws.max_column))
    print("=" * 110)
    for r in rows:
        cells = []
        for c in cols:
            v = ws.cell(r, c).value
            if v is None:
                continue
            if isinstance(v, float):
                v = "%.2f" % v
            cells.append("%s%s=%s" % (cl(c), r, v))
        if cells:
            print("  " + " | ".join(cells))
    merges = [str(m) for m in ws.merged_cells.ranges]
    print("  MERGED (%d): %s" % (len(merges), ", ".join(sorted(merges)[:80])))
    print()


# ---- CDJ: month labels live in col C? Are they merged/single-cell? ----
dump("CDJ", [11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 30, 31, 32, 40, 41, 42, 50, 51, 52, 60, 61, 62],
     list(range(1, 23)))
# ---- CRJ ----
dump("CRJ", [8, 9, 10, 24, 25, 26, 35, 36, 37, 38, 60, 61, 62, 100, 101, 102, 700, 701, 702, 710, 711],
     list(range(1, 17)))
# ---- GJ ----
dump("GJ", [9, 10, 11, 12, 13, 14, 15, 16, 20, 21, 22, 30, 31, 32, 40, 41, 42, 50, 51, 52, 60, 61, 62, 70, 71],
     list(range(1, 9)))
# ---- GL header band + block anchors + a few month rows each ----
dump("GL", [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 24, 25, 26, 208, 221, 598, 608, 609, 610, 611],
     list(range(1, 12)))
# ---- GL_V8_CALC ----
dump("GL_V8_CALC", [1, 2, 3, 4, 5, 6, 555, 556, 557], list(range(1, 7)))
# ---- GL_AUDIT ----
dump("GL_AUDIT", list(range(1, 20)), list(range(1, 5)))
wb.close()
