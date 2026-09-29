"""review_v82.py - read-only review of Global-Smile_2026-v8.2-final-review.xlsm.

Answers: can RefreshGL refresh ALL 46 COAs for EVERY month?

Read-only: never writes to the workbook binary.
Outputs: v8.2_review/  (VBA extraction + JSON/text evidence)
"""
import hashlib
import io
import json
import os
import re
import zipfile

import openpyxl
from oletools.olevba import VBA_Parser

BASE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(BASE, "Global-Smile_2026-v8.2-final-review.xlsm")
OUT = os.path.join(BASE, "v8.2_review")


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def norm(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\u00a0", " ")).strip().upper()


def main():
    os.makedirs(OUT, exist_ok=True)
    report = {}
    report["workbook"] = os.path.basename(WB)
    report["workbook_sha256"] = sha256_file(WB)

    # ---- VBA project ----
    with zipfile.ZipFile(WB) as z:
        names = z.namelist()
        vp = z.read("xl/vbaProject.bin") if "xl/vbaProject.bin" in names else None
    report["vbaProject_bin_sha256"] = hashlib.sha256(vp).hexdigest() if vp else None

    mods = []
    parser = VBA_Parser(WB)
    vba_text = {}
    try:
        for (_, stream, fname, code) in parser.extract_macros():
            if not code.strip():
                continue
            mods.append({"name": fname, "chars": len(code),
                         "lines": code.count("\n") + (0 if code.endswith("\n") else 1)})
            vba_text[fname] = code
    finally:
        parser.close()
    report["modules"] = mods
    with open(os.path.join(OUT, "v8.2_vba_dump.txt"), "w", encoding="utf-8", newline="") as fh:
        for k, v in vba_text.items():
            fh.write("=" * 78 + f"\n== MODULE: {k}\n" + "=" * 78 + "\n" + v + "\n\n")

    wb = openpyxl.load_workbook(WB, keep_vba=True, data_only=True)
    report["sheets"] = wb.sheetnames

    # ---- GL sheet structure: 46 blocks, stride 13, row 13 start, col F title ----
    gl = wb["GL"]
    blocks = []
    for b in range(46):
        r0 = 13 + b * 13
        title = gl.cell(r0, 6).value
        months = []
        for m in range(12):
            r = r0 + m
            months.append({
                "row": r,
                "B": gl.cell(r, 2).value,
                "G": gl.cell(r, 7).value,
                "H": gl.cell(r, 8).value,
            })
        blocks.append({"block": b, "anchor_row": r0, "title": title,
                       "title_norm": norm(title), "months": months})
    report["gl_blocks"] = blocks
    report["gl_max_row"] = gl.max_row
    report["gl_max_col"] = gl.max_column

    # Non-empty rows beyond block 45 (row 13+45*13+11 = 609)
    extra = []
    for r in range(610, gl.max_row + 1):
        vals = [(c, gl.cell(r, c).value) for c in range(1, 12) if gl.cell(r, c).value not in (None, "")]
        if vals:
            extra.append({"row": r, "cells": vals})
    report["gl_rows_below_block45"] = extra

    # ---- GL titles list ----
    titles = [b["title"] for b in blocks]
    report["gl_title_count"] = len(titles)
    report["gl_titles"] = titles
    dupes = sorted({t for t in titles if titles.count(t) > 1}, key=str)
    report["gl_title_duplicates"] = dupes

    for name in ("CDJ", "CRJ", "GJ"):
        ws = wb[name]
        report[f"{name}_max_row"] = ws.max_row
        report[f"{name}_max_col"] = ws.max_column

    wb.close()
    with open(os.path.join(OUT, "review_v82_structure.json"), "w", encoding="utf-8") as fh:
        json.dump(report, fh, indent=2, default=str)

    print("workbook sha256      :", report["workbook_sha256"])
    print("vbaProject.bin sha256:", report["vbaProject_bin_sha256"])
    print("sheets               :", report["sheets"])
    print("modules              :", len(mods))
    for m in mods:
        print("   %-35s lines=%s" % (m["name"], m["lines"]))
    print("GL titles (%d):" % len(titles))
    for i, b in enumerate(blocks):
        print("   %2d row %3d  %r" % (i, b["anchor_row"], b["title"]))
    print("duplicates:", dupes)
    print("rows below block45:", extra)


if __name__ == "__main__":
    main()
