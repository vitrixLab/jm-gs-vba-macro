"""scan_cdcrj.py — read-only scan of CDJ + CRJ posting populations for the
RefreshGL-into-GL build. Prints per-account x month debit/credit from CDJ F:S
(signed) + CRJ H debit / I:L credits, and the expected GL row writes.
"""
import openpyxl
from collections import defaultdict

WB = "Global-Smile_2026-v8.0.xlsm"
wb = openpyxl.load_workbook(WB, data_only=True)

MONTHS = {"JAN": 1, "JANUARY": 1, "FEB": 2, "FEBRUARY": 2, "MAR": 3, "MARCH": 3,
          "APR": 4, "APRIL": 4, "MAY": 5, "JUN": 6, "JUNE": 6, "JUL": 7, "JULY": 7,
          "AUG": 8, "AUGUST": 8, "SEP": 9, "SEPTEMBER": 9, "OCT": 10, "OCTOBER": 10,
          "NOV": 11, "NOVEMBER": 11, "DEC": 12, "DECEMBER": 12}

CDJMAP = {6: "Cash in Bank", 7: "Input VAT", 8: "EWT Payable", 9: "Petty Cash Fund",
          10: "Government Contributions (EE)", 11: "Government Loans (EE)",
          12: "De Minimis", 13: "Salaries and Wages", 14: "Clinic Material and Supplies",
          15: "Rent", 16: "Gas, Oil, Parking, Toll Fees", 18: "Professional Fees"}
CRJMAP = {9: "Excess CWT Over IT", 10: "VAT Payable", 11: "Sales", 12: "Sales"}

mat = defaultdict(lambda: [0.0, 0.0])
gd = gc = 0.0
inv = 0

# CDJ
ws = wb["CDJ"]
cur = 0
last = ws.max_row
for r in range(15, last + 1):
    mv = ws.cell(r, 3).value
    if isinstance(mv, str) and mv.strip().upper() in MONTHS:
        cur = MONTHS[mv.strip().upper()]
    pv = ws.cell(r, 5).value
    if not cur or not pv or not str(pv).strip() or str(pv).strip().upper() == "TOTAL":
        continue
    for c in range(6, 20):
        v = ws.cell(r, c).value
        if isinstance(v, str):
            continue
        if v is None:
            continue
        if not isinstance(v, (int, float)):
            inv += 1
            continue
        if abs(v) < 0.005:
            continue
        a = CDJMAP.get(c)
        if not a:
            continue  # Q/S ignored (verified zero)
        if v >= 0:
            mat[(a, cur)][0] += v
            gd += v
        else:
            mat[(a, cur)][1] += -v
            gc += -v

# CRJ
ws = wb["CRJ"]
for r in range(10, ws.max_row + 1):
    f = ws.cell(r, 7).value
    if not f or not str(f).strip() or str(f).strip().upper() == "TOTAL":
        continue
    log = ws.cell(r, 15).value
    dt = None
    if hasattr(log, "year"):
        dt = log
    else:
        s = str(log or "")
        if " | " in s:
            s = s.rsplit(" | ", 1)[-1]
        try:
            from datetime import datetime as _dt
            dt = _dt.strptime(s.strip(), "%m/%d/%Y %I:%M:%S %p")
        except Exception:
            dt = None
    if dt is None or dt.year != 2026:
        continue
    m = dt.month
    v = ws.cell(r, 8).value
    if isinstance(v, (int, float)) and abs(v) >= 0.005:
        mat[("Cash in Bank", m)][0] += v
        gd += v
    for c in range(9, 13):
        v = ws.cell(r, c).value
        if isinstance(v, (int, float)) and abs(v) >= 0.005:
            mat[(CRJMAP[c], m)][1] += v
            gc += v

print("CDJ+CRJ gross debit=%.2f credit=%.2f diff=%.2f invalid=%d" % (gd, gc, gd - gc, inv))
print("nonzero account-months:", len(mat))
accts = sorted(set(a for (a, m) in mat))
print("accounts touched:", len(accts))
for a in accts:
    cells = ["%s:%d=%.2f/%.2f" % (a, m, mat[(a, m)][0], mat[(a, m)][1]) for m in range(1, 13) if mat[(a, m)][0] or mat[(a, m)][1]]
    print("  " + "; ".join(cells))
