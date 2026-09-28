"""Full GL formula inventory + drawing macro scan for v8.0."""
import glob
import os
import re
import zipfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(HERE, "Global-Smile_2026-v8.0.xlsm")

wb = load_workbook(WB, data_only=False)
ws = wb["GL"]

print("== GL rows with any content in B..I (formulas view) ==")
last_row = 0
for r in range(11, 620):
    vals = []
    has_content = False
    for c in range(2, 10):
        v = ws.cell(row=r, column=c).value
        if v is not None and str(v).strip():
            has_content = True
            vals.append(f"{chr(64+c)}={v}")
    if has_content:
        last_row = r
        print(f"r{r}: " + " | ".join(vals))
print(f"last GL content row: {last_row}")

print()
print("== Drawing / macro button assignments ==")
with zipfile.ZipFile(WB) as z:
    for n in z.namelist():
        if ("drawing" in n and n.endswith(".xml")) or n.endswith(".vml"):
            data = z.read(n).decode("utf-8", "replace")
            macros = re.findall(r'macro="([^"]+)"|&amp;macroname="([^"]+)"|procedural="([^"]+)"', data)
            flat = [m for tup in macros for m in tup if m]
            if flat:
                print(f"{n}: {flat}")
        if n == "xl/workbook.xml":
            data = z.read(n).decode("utf-8", "replace")
            print("definedNames:", re.findall(r'<definedName[^>]*name="([^"]+)"[^>]*>([^<]*)</definedName>', data)[:40])
