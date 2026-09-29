"""inspect_v82_sheets.py - dump CDJ/CRJ/GJ/GL layout of the v8.2 workbook (read-only)."""
import os
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(BASE, "Global-Smile_2026-v8.2-final-review.xlsm")
wb = openpyxl.load_workbook(WB, keep_vba=True, data_only=True)


def cl(i):
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def dump(name, rows, cols, sheet=None):
    ws = wb[sheet or name]
    print("=" * 100)
    print("SHEET %s  dims=%s  max_row=%s max_col=%s" % (name, ws.dimensions, ws.max_row, ws.max_column))
    print("=" * 100)
    for r in rows:
        cells = []
        for c in cols:
            v = ws.cell(r, c).value
            if v is None:
                continue
            if isinstance(v, float):
                v = ("%.2f" % v)
            cells.append("%s%s=%s" % (cl(c), r, v))
        if cells:
            print("  " + " | ".join(cells))
    print()


# CDJ: headers + first data rows + month/total rows in col C / col E
dump("CDJ", list(range(1, 26)) + list(range(26, 40)), list(range(1, 23)))
# CRJ
dump("CRJ", list(range(1, 26)) + list(range(26, 40)), list(range(1, 17)))
# GJ
dump("GJ", list(range(1, 30)) + list(range(30, 45)), list(range(1, 9)))
# GL top area
dump("GL", list(range(1, 30)), list(range(1, 12)))
# GL_V8_CALC head
dump("GL_V8_CALC", list(range(1, 12)), list(range(1, 7)))
# GL_AUDIT
dump("GL_AUDIT", list(range(1, 30)), list(range(1, 5)))
# COA_MAP
dump("COA_MAP", list(range(1, 15)), list(range(1, 7)))
# V8.2_FINAL_REVIEW
dump("V8.2_FINAL_REVIEW", list(range(1, 60)), list(range(1, 7)))
wb.close()
