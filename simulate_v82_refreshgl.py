"""simulate_v82_refreshgl.py - replicate modEngine.RefreshGL (as embedded in
Global-Smile_2026-v8.2-final-review.xlsm) over the workbook's own CDJ/CRJ/GJ data.

Read-only. Answers: does RefreshGL produce a value for all 46 COAs x 12 months,
and are those values right / complete?
"""
import json
import os
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(BASE, "Global-Smile_2026-v8.2-final-review.xlsm")
OUT = os.path.join(BASE, "v8.2_review")

wb = openpyxl.load_workbook(WB, keep_vba=True, data_only=True)


def GLN(s):
    if s is None:
        return ""
    s = str(s).strip().lower().replace("'", "")
    while "  " in s:
        s = s.replace("  ", " ")
    return s


def GLMonth(v):
    k = GLN(v)
    m = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
         "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}
    if k[:3] in m:
        return m[k[:3]]
    t = str(v).strip()
    mo = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$", t)
    if mo:
        return int(mo.group(1))
    return 0


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


GLA_MAP = {
    "cash in bank": "Cash in Bank", "cash": "Cash in Bank",
    "petty cash fund": "Petty Cash Fund", "petty cash": "Petty Cash Fund",
    "account receivables": "Account Receivables", "accounts receivable": "Account Receivables",
    "account receivable": "Account Receivables",
    "advances to employee": "Advances to Employees", "advances to employees": "Advances to Employees",
    "input vat": "Input VAT", "input tax": "Input VAT",
    "excess cwt over it": "Excess CWT Over IT", "cwt/it": "Excess CWT Over IT", "cwt": "Excess CWT Over IT",
    "leasehold improvements": "Leasehold Improvements",
    "dental equipment": "Dental Equipment", "medical equipment": "Dental Equipment",
    "due to clinicians": "Due To Clinicians",
    "due to clinicians- visiting": "Due to Clinicians- Visiting",
    "due to clinicians visiting": "Due to Clinicians- Visiting",
    "vat payable": "VAT Payable", "output vat payable": "VAT Payable", "output vat": "VAT Payable",
    "ewt payable": "EWT Payable", "withholding tax": "EWT Payable",
    "withholding tax payable": "EWT Payable", "ewt": "EWT Payable",
    "government contributions (ee)": "Government Contributions (EE)",
    "ee government contributions": "Government Contributions (EE)",
    "government loans (ee)": "Government Loans (EE)", "ee government loans": "Government Loans (EE)",
    "income tax payable": "Income Tax Payable",
    "paid-up capital": "Paid-Up Capital", "paid up capital": "Paid-Up Capital",
    "retained earnings (deficit)": "Retained Earnings (Deficit)",
    "retained earnings": "Retained Earnings (Deficit)",
    "sales": "Sales", "exempt sales": "Sales",
    "rent": "Rent", "rental": "Rent",
    "clinicians fee-corporators": "Clinician's Fee-Corporators",
    "clinicians fee-visiting": "Clinician's Fee-Visiting",
    "clinic material and supplies": "Clinic Material and Supplies",
    "clinic materials and supplies": "Clinic Material and Supplies",
    "clinic supplies": "Clinic Material and Supplies",
    "common use service area": "Common Use Service Area",
    "light and water expense": "Light and Water Expense", "light and water": "Light and Water Expense",
    "depn expense-leasehold improvements": "Dep'n Expense-Leasehold Improvements",
    "depn expense-dental equipment": "Dep'n Expense-Dental Equipment",
    "salaries and wages": "Salaries and Wages", "salary and wages": "Salaries and Wages",
    "de minimis": "De Minimis",
    "government contributions er share": "Government Contributions ER Share",
    "government contributions er": "Government Contributions ER Share",
    "professional fees": "Professional Fees", "professional fee": "Professional Fees",
    "repair and maintenance": "Repair and Maintenance", "repairs and maintenance": "Repair and Maintenance",
    "pantry supplies": "Pantry Supplies",
    "office supplies": "Office Supplies", "supplies": "Office Supplies",
    "meals, gifts & entertainment expense": "Meals, Gifts & Entertainment Expense",
    "meals and entertainment": "Meals, Gifts & Entertainment Expense",
    "communication": "Communication ",
    "printing and duplication": "Printing and Duplication",
    "taxes and licences": "Taxes and Licences", "taxes and licenses": "Taxes and Licences",
    "delivery": "Delivery",
    "software subscription": "Software Subscription",
    "miscellaneous": "Miscellaneous", "charges": "Miscellaneous", "bank charge": "Miscellaneous",
    "transportation and travel": "Transportation and Travel",
    "transporation and travel": "Transportation and Travel",
    "advertisement": "Advertisement",
    "marketing expense": "Marketing Expense",
    "gas, oil, parking, toll fees": "Gas, Oil, Parking, Toll Fees",
    "fuel and oil": "Gas, Oil, Parking, Toll Fees",
    "other penalties and charges": "Other Penalties and Charges",
    "insurance": "Insurance",
}


def GLAcct(s):
    return GLA_MAP.get(GLN(s), "")


class D(dict):
    def add(self, a, m, dr, cr):
        if a == "" or m < 1 or m > 12:
            return
        k = a + "|" + str(m)
        x = self.get(k, [0.0, 0.0])
        self[k] = [x[0] + dr, x[1] + cr]


GL_ACCOUNTS = [
    "Cash in Bank", "Petty Cash Fund", "Account Receivables", "Advances to Employees",
    "Input VAT", "Excess CWT Over IT", "Leasehold Improvements", "Dental Equipment",
    "Due To Clinicians", "Due to Clinicians- Visiting", "VAT Payable", "EWT Payable",
    "Government Contributions (EE)", "Government Loans (EE)", "Income Tax Payable",
    "Paid-Up Capital", "Retained Earnings (Deficit)", "Sales", "Rent",
    "Clinician's Fee-Corporators", "Clinician's Fee-Visiting", "Clinic Material and Supplies",
    "Common Use Service Area", "Light and Water Expense",
    "Dep'n Expense-Leasehold Improvements", "Dep'n Expense-Dental Equipment",
    "Salaries and Wages", "De Minimis", "Government Contributions ER Share",
    "Professional Fees", "Repair and Maintenance", "Pantry Supplies", "Office Supplies",
    "Meals, Gifts & Entertainment Expense", "Communication ", "Printing and Duplication",
    "Taxes and Licences", "Delivery", "Software Subscription", "Miscellaneous",
    "Transportation and Travel", "Advertisement", "Marketing Expense",
    "Gas, Oil, Parking, Toll Fees", "Other Penalties and Charges", "Insurance",
]

ev = {"cdj_skipped_no_month": [], "cdj_rows": 0, "crj_blocks": [], "crj_rows": [],
      "gj_rows": 0, "audit": [], "cdj_cash_out_mismatch": []}


def scan_cdj(d):
    w = wb["CDJ"]
    for r in range(15, w.max_row + 1):
        if GLN(w.cell(r, 5).value) != "total":
            m = GLMonth(w.cell(r, 3).value)
            if m > 0:
                ev["cdj_rows"] += 1
                v = num(w.cell(r, 6).value)
                if v is not None:
                    if v < 0:
                        d.add("Cash in Bank", m, 0, -v)
                    else:
                        d.add("Cash in Bank", m, v, 0)
                for col, acct in ((7, "Input VAT"), (8, "EWT Payable"), (9, "Petty Cash Fund"),
                                  (10, "Government Contributions (EE)"),
                                  (11, "Government Loans (EE)"), (12, "De Minimis"),
                                  (13, "Salaries and Wages"), (14, "Clinic Material and Supplies"),
                                  (15, "Rent"), (16, "Gas, Oil, Parking, Toll Fees"),
                                  (17, "Communication "), (18, "Professional Fees")):
                    v = num(w.cell(r, col).value)
                    if v is not None:
                        if col in (8, 10, 11):
                            d.add(acct, m, 0, v)
                        else:
                            d.add(acct, m, v, 0)
                a = GLAcct(w.cell(r, 19).value)
                if a != "":
                    v = num(w.cell(r, 20).value)
                    if v is not None:
                        d.add(a, m, v, 0)
                    v = num(w.cell(r, 21).value)
                    if v is not None:
                        d.add(a, m, 0, v)
                elif str(w.cell(r, 19).value or "").strip() != "":
                    ev["audit"].append(["CDJ", str(w.cell(r, 19).value), 0])
            else:
                vals = [num(w.cell(r, c).value) for c in range(6, 22)]
                lab = str(w.cell(r, 19).value or "").strip()
                if any(x is not None for x in vals) or lab:
                    ev["cdj_skipped_no_month"].append(
                        {"row": r, "C": w.cell(r, 3).value, "E": w.cell(r, 5).value,
                         "S": lab, "T": w.cell(r, 20).value, "U": w.cell(r, 21).value,
                         "F": w.cell(r, 6).value})


def scan_crj(d):
    w = wb["CRJ"]
    ends = []
    for r in range(10, w.max_row + 1):
        if GLN(w.cell(r, 6).value) == "total" and GLMonth(w.cell(r, 4).value) > 0:
            ends.append(r)
    lastf = None
    for r in range(10, w.max_row + 1):
        if str(w.cell(r, 6).value or "").strip() != "":
            lastf = r
    st = 10
    for m in range(1, 13):
        en = (ends[m - 1] - 1) if m <= len(ends) else lastf
        for rr in range(st, (en or 0) + 1):
            if GLN(w.cell(rr, 6).value) != "total":
                for col, acct, side in ((8, "Cash in Bank", "dr"), (9, "Excess CWT Over IT", "dr"),
                                        (10, "VAT Payable", "cr"), (11, "Sales", "cr"),
                                        (12, "Sales", "cr")):
                    v = num(w.cell(rr, col).value)
                    if v is not None:
                        ev["crj_rows"].append({"gl_month": m, "row": rr, "col": col,
                                               "acct": acct, "amt": v,
                                               "log": w.cell(rr, 15).value})
                        d.add(acct, m, v if side == "dr" else 0, v if side == "cr" else 0)
        if m <= len(ends):
            st = ends[m - 1] + 1
        else:
            break
    ev["crj_blocks"] = [{"row": e, "label": w.cell(e, 4).value} for e in ends]


def scan_gj(d):
    w = wb["GJ"]
    m = 0
    for r in range(10, w.max_row + 1):
        mm = GLMonth(w.cell(r, 2).value)
        if mm > 0:
            m = mm
        s = str(w.cell(r, 5).value or "").strip()
        if m > 0 and s != "" and GLN(s) != "total":
            ev["gj_rows"] += 1
            a = GLAcct(s)
            dr = num(w.cell(r, 6).value) or 0
            cr = num(w.cell(r, 7).value) or 0
            if a != "":
                d.add(a, m, dr, cr)
            elif dr != 0 or cr != 0:
                ev["audit"].append(["GJ", s, dr + cr])


d = D()
scan_cdj(d)
scan_crj(d)
scan_gj(d)

gd = sum(v[0] for v in d.values())
gc = sum(v[1] for v in d.values())

grid = {}
for a in GL_ACCOUNTS:
    for m in range(1, 13):
        grid[(a, m)] = d.get(a + "|" + str(m), [0.0, 0.0])

zero_cells = [(a, m) for (a, m), x in grid.items() if abs(x[0]) < 1e-9 and abs(x[1]) < 1e-9]
zero_accounts = [a for a in GL_ACCOUNTS
                 if all(abs(grid[(a, m)][0]) < 1e-9 and abs(grid[(a, m)][1]) < 1e-9
                        for m in range(1, 13))]
month_tot = {m: (sum(grid[(a, m)][0] for a in GL_ACCOUNTS),
                 sum(grid[(a, m)][1] for a in GL_ACCOUNTS)) for m in range(1, 13)}
MN = ["", "JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]

print("=" * 100)
print("SIMULATION OF modEngine.RefreshGL (as embedded in the v8.2 workbook)")
print("=" * 100)
print("total debit  = %.2f" % gd)
print("total credit = %.2f" % gc)
print("difference   = %.2f" % (gd - gc))
print("accounts written (GLAccounts array) = %d" % len(GL_ACCOUNTS))
print("rows written = %d  (GL rows 13..%d, contiguous, no separators)"
      % (len(GL_ACCOUNTS) * 12, 13 + len(GL_ACCOUNTS) * 12 - 1))
print("zero cells (0 debit and 0 credit) = %d of %d" % (len(zero_cells), len(grid)))
print("accounts with NO activity in ANY month (%d): %s" % (len(zero_accounts), zero_accounts))
print()
print("per-month totals written by RefreshGL:")
for m in range(1, 13):
    print("  %-3s debit=%12.2f  credit=%12.2f" % (MN[m], month_tot[m][0], month_tot[m][1]))
print()
print("CDJ rows posted = %d ; CDJ data rows SILENTLY SKIPPED (no month in col C) = %d"
      % (ev["cdj_rows"], len(ev["cdj_skipped_no_month"])))
for s in ev["cdj_skipped_no_month"]:
    print("   SKIP CDJ row %-4s F=%-10s S=%-14s T=%-9s U=%-6s %s"
          % (s["row"], s["F"], s["S"], s["T"], s["U"], (s["E"] or "")[:30]))
print()
print("CRJ TOTAL blocks found = %d : %s" % (len(ev["crj_blocks"]), ev["crj_blocks"]))
print("CRJ data rows posted = %d" % len(ev["crj_rows"]))
for x in ev["crj_rows"]:
    print("   CRJ row %-4s booked to month %-2s col=%s %s %.2f   (entry log: %s)"
          % (x["row"], x["gl_month"], x["col"], x["acct"], x["amt"], x["log"]))
print()
print("GJ rows posted = %d" % ev["gj_rows"])
print("GL_AUDIT unmapped rows the engine would write (%d):" % len(ev["audit"]))
for a in ev["audit"]:
    print("   %-4s %-40s %s" % (a[0], a[1], a[2]))

with open(os.path.join(OUT, "refreshgl_simulation.json"), "w", encoding="utf-8") as fh:
    json.dump({"total_debit": gd, "total_credit": gc, "difference": gd - gc,
               "accounts": len(GL_ACCOUNTS), "rows_written": len(GL_ACCOUNTS) * 12,
               "zero_cells": len(zero_cells), "total_cells": len(grid),
               "accounts_with_no_activity": zero_accounts,
               "month_totals": month_tot,
               "grid": {"%s|%d" % (a, m): grid[(a, m)] for (a, m) in grid},
               "evidence": ev}, fh, indent=2, default=str)
wb.close()

