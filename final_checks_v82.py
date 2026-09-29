"""final_checks_v82.py - GL current values + stray CRJ rows (read-only)."""
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(BASE, "Global-Smile_2026-v8.2-final-review.xlsm")
wb = openpyxl.load_workbook(WB, keep_vba=True, data_only=True)

gl = wb["GL"]
td = tc = 0.0
nz = []
for r in range(13, 610):
    g, h = gl.cell(r, 7).value, gl.cell(r, 8).value
    if isinstance(g, (int, float)):
        td += g
    if isinstance(h, (int, float)):
        tc += h
    if (isinstance(g, (int, float)) and g) or (isinstance(h, (int, float)) and h):
        nz.append((r, gl.cell(r, 6).value, gl.cell(r, 2).value, g, h, gl.cell(r, 9).value))
print("GL rows 13..609: total G(debit)=%.2f  total H(credit)=%.2f  difference=%.2f" % (td, tc, td - tc))
print("non-zero GL cells: %d" % len(nz))
for x in nz:
    print("   row %-4s %-20s %-5s G=%-10s H=%-10s I(ending)=%s" % x)
print()
print("GL col A (row label) sample:", [(r, gl.cell(r, 1).value) for r in (11, 12, 13, 14, 25, 26)])
print("GL row 610..620:", [(r, gl.cell(r, 6).value) for r in range(610, 621)])
print("GL formulas? sample G13 type:", type(gl.cell(13, 7).value).__name__)
wbf = openpyxl.load_workbook(WB, keep_vba=True, data_only=False)
glf = wbf["GL"]
print("formula sample F13/G13/I13:", repr(glf.cell(13, 6).value), repr(glf.cell(13, 7).value), repr(glf.cell(13, 9).value))
print("GL merged/ranges that reference F13:F609? checked separately")
print()
crj = wb["CRJ"]
print("CRJ rows 575-590:")
for r in range(575, 591):
    vals = [(c, crj.cell(r, c).value) for c in range(4, 16) if crj.cell(r, c).value not in (None, "")]
    if vals:
        print("   row %-4s %s" % (r, vals))
print("CRJ rows 26-35 (FEB block, should be empty data):")
for r in range(26, 36):
    vals = [(c, crj.cell(r, c).value) for c in range(4, 16) if crj.cell(r, c).value not in (None, "")]
    print("   row %-4s %s" % (r, vals))
print("CRJ rows 710-719:")
for r in range(710, 720):
    vals = [(c, crj.cell(r, c).value) for c in range(4, 16) if crj.cell(r, c).value not in (None, "")]
    if vals:
        print("   row %-4s %s" % (r, vals))
wb.close()
wbf.close()
