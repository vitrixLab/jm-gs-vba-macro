"""Verify CDJ cols 9-13 headers, month markers, CRJ sample rows."""
import openpyxl

wb = openpyxl.load_workbook(r"D:\citrixlabph\globalsmile\Global-Smile_2026-v8.0.xlsm", data_only=True)
cd = wb["CDJ"]

print("== CDJ headers H13..N14 (cols 8-14) ==")
for r in (13, 14):
    print(f"r{r}:", [(chr(64+c), cd.cell(row=r, column=c).value) for c in range(8, 15)])

print()
print("== CDJ C/D values rows 13-20 (month markers, day) ==")
for r in range(13, 21):
    print(f"r{r}: C={cd.cell(row=r, column=3).value!r} D={cd.cell(row=r, column=4).value!r} E={cd.cell(row=r, column=5).value!r} X={cd.cell(row=r, column=24).value!r}")

print()
print("== CDJ last data row per col C/E ==")
last = cd.Cells if False else None
print("max_row:", cd.max_row)
for r in (57, 58, 59):
    print(f"r{r}: C={cd.cell(row=r, column=3).value!r} E={cd.cell(row=r, column=5).value!r}")

print()
print("== CRJ sample data rows 10-12 (D..Q) ==")
cr = wb["CRJ"]
for r in (10, 11, 12):
    print(f"r{r}:", [cr.cell(row=r, column=c).value for c in range(4, 18)])

print()
print("== GL rows 11-12 headers ==")
gl = wb["GL"]
for r in (11, 12):
    print(f"r{r}:", [gl.cell(row=r, column=c).value for c in range(2, 10)])
