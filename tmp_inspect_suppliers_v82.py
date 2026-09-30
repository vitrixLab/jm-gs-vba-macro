import openpyxl
SRC = 'Global-Smile_2026-v8.2-final-review.xlsm'
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True, keep_vba=True)
print('sheets=', wb.sheetnames)
for name in ['SUPPLIERS DATA', 'COA_MAP', 'PJ', 'CRJ', 'SJ']:
    if name not in wb.sheetnames:
        print(name, 'MISSING')
        continue
    ws = wb[name]
    print('\n====', name, 'max=', ws.max_row, ws.max_column)
    hi = min(ws.max_row, 8)
    for r in range(1, hi + 1):
        print(r, [ws.cell(r, c).value for c in range(1, 11)])
    if ws.max_row > 8:
        for r in [9, 10, ws.max_row - 1, ws.max_row]:
            print(r, [ws.cell(r, c).value for c in range(1, 11)])
wb.close()
