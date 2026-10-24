import openpyxl

wb = openpyxl.load_workbook(r'D:\citrixlabph\globalsmile_v2\Global-Smile_2026-v8.4-NO-SHIT-FIXED.xlsm', data_only=True)
gl = wb['GL']
cdj = wb['CDJ']
crj = wb['CRJ']
audit = wb['GL_AUDIT']

print("=== 1. GL_AUDIT VERIFICATION ===")
last_audit_row = audit.max_row
while last_audit_row >= 1 and all(audit.cell(row=last_audit_row, column=c).value is None for c in range(1, 5)):
    last_audit_row -= 1

print(f"Audit Row {last_audit_row}:")
print(f"  Timestamp:  {audit.cell(row=last_audit_row, column=1).value}")
print(f"  Status:     {audit.cell(row=last_audit_row, column=2).value}")
print(f"  Summary:    {audit.cell(row=last_audit_row, column=3).value}")

# Read all 46 GL accounts
MONTH_NAMES = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]

accounts_data = []
monthly_totals_dr = [0.0] * 12
monthly_totals_cr = [0.0] * 12

for b in range(46):
    r_start = 13 + b * 13
    acct_title = str(gl.cell(row=r_start, column=6).value or "").strip()
    months = []
    prev_bal = 0.0
    rollforward_ok = True
    
    for m_idx in range(12):
        r = r_start + m_idx
        m_name = gl.cell(row=r, column=2).value
        dr = float(gl.cell(row=r, column=7).value or 0.0)
        cr = float(gl.cell(row=r, column=8).value or 0.0)
        bal = float(gl.cell(row=r, column=9).value or 0.0)
        
        monthly_totals_dr[m_idx] += dr
        monthly_totals_cr[m_idx] += cr
        
        expected_bal = prev_bal + dr - cr
        if abs(bal - expected_bal) > 0.001:
            rollforward_ok = False
        prev_bal = bal
        
        months.append((m_name, dr, cr, bal))
        
    accounts_data.append((b + 1, r_start, acct_title, rollforward_ok, months))

print("\n=== 2. MONTHLY TOTALS ACROSS ALL 46 ACCOUNTS ===")
total_all_dr = sum(monthly_totals_dr)
total_all_cr = sum(monthly_totals_cr)
print(f"{'Month':<6} | {'Total Debit':>15} | {'Total Credit':>15} | {'Diff (Dr-Cr)':>12}")
print("-" * 55)
for m_idx in range(12):
    m_name = MONTH_NAMES[m_idx]
    dr = monthly_totals_dr[m_idx]
    cr = monthly_totals_cr[m_idx]
    diff = dr - cr
    print(f"{m_name:<6} | {dr:>15,.2f} | {cr:>15,.2f} | {diff:>12.2f}")
print("-" * 55)
print(f"{'TOTAL':<6} | {total_all_dr:>15,.2f} | {total_all_cr:>15,.2f} | {total_all_dr - total_all_cr:>12.2f}")

print("\n=== 3. JANUARY (MONTH 1) ACCOUNT-LEVEL BREAKDOWN ===")
print(f"{'Blk':>3} | {'Account Title':<38} | {'Jan Debit':>14} | {'Jan Credit':>14} | {'Jan Ending Bal':>15}")
print("-" * 92)
jan_active = 0
for b_num, r_start, acct_title, rollforward_ok, months in accounts_data:
    m_name, dr, cr, bal = months[0]
    if abs(dr) > 0.001 or abs(cr) > 0.001 or abs(bal) > 0.001:
        jan_active += 1
        print(f"{b_num:3d} | {acct_title:<38} | {dr:>14,.2f} | {cr:>14,.2f} | {bal:>15,.2f}")
print("-" * 92)
print(f"Total active accounts in Jan: {jan_active} of 46")
print(f"Jan Total Debit:  {monthly_totals_dr[0]:>15,.2f}")
print(f"Jan Total Credit: {monthly_totals_cr[0]:>15,.2f}")
print(f"Jan Difference:   {monthly_totals_dr[0] - monthly_totals_cr[0]:>15.2f}")

print("\n=== 4. ROLL-FORWARD INTEGRITY CHECK ===")
failed_rollforwards = []
for b_num, r_start, acct_title, rollforward_ok, months in accounts_data:
    if not rollforward_ok:
        failed_rollforwards.append((b_num, acct_title))
if not failed_rollforwards:
    print("ALL 46 ACCOUNTS PASS ROLL-FORWARD INTEGRITY: Ending Balance = Prev Balance + Debit - Credit for all 12 months.")
else:
    print(f"WARNING: {len(failed_rollforwards)} accounts failed roll-forward:")
    for b_num, acct_title in failed_rollforwards:
        print(f"  Block {b_num}: {acct_title}")

print("\n=== 5. ALL 46 ACCOUNTS SUMMARY (12 MONTHS) ===")
print(f"{'Blk':>3} | {'Account Title':<38} | {'Total Debit':>14} | {'Total Credit':>14} | {'Dec Ending Bal':>15}")
print("-" * 92)
active_total = 0
for b_num, r_start, acct_title, rollforward_ok, months in accounts_data:
    tot_dr = sum(m[1] for m in months)
    tot_cr = sum(m[2] for m in months)
    dec_bal = months[11][3]
    if abs(tot_dr) > 0.001 or abs(tot_cr) > 0.001:
        active_total += 1
        print(f"{b_num:3d} | {acct_title:<38} | {tot_dr:>14,.2f} | {tot_cr:>14,.2f} | {dec_bal:>15,.2f}")
print("-" * 92)
print(f"Active accounts across the year: {active_total} of 46")
