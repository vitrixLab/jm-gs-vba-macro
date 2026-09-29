"""stray_v82.py - leftover cells outside the GL/CRJ written ranges (read-only)."""
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
WB = os.path.join(BASE, "Global-Smile_2026-v8.2-final-review.xlsm")
wb = openpyxl.load_workbook(WB, keep_vba=True, data_only=True)

gl = wb["GL"]
print("--- GL columns J..AA non-empty cells (rows 1..1000) ---")
cnt = 0
for r in range(1, 1001):
    vals = [(c, gl.cell(r, c).value) for c in range(10, 28) if gl.cell(r, c).value not in (None, "")]
    if vals:
        cnt += 1
        if cnt <= 25:
            print("   row %-4s %s" % (r, vals))
print("   total rows with J..AA content: %d" % cnt)
print("   GL print area:", gl.print_area, "| autofilter:", gl.auto_filter.ref)

crj = wb["CRJ"]
print()
print("--- CRJ rows 570-590, cols F..P ---")
for r in range(570, 591):
    vals = [(c, crj.cell(r, c).value) for c in range(6, 17) if crj.cell(r, c).value not in (None, "")]
    if vals:
        print("   row %-4s %s" % (r, vals))
print()
print("--- CRJ rows 110-125 (JAN-DEC block ends) ---")
for r in range(110, 126):
    vals = [(c, crj.cell(r, c).value) for c in range(4, 17) if crj.cell(r, c).value not in (None, "")]
    if vals:
        print("   row %-4s %s" % (r, vals))
wb.close()
